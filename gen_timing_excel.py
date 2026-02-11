import load_timing_excel as ld
import simulate as sim
import draw_timing as dr
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from collections import defaultdict


EMPTY_ROW_HEIGHT = 5

def split_hier(sig):
    if "__" in sig:
        inst, local = sig.split("__", 1)
        return inst, local
    return None, sig

def set_wave_column_width(ws, col_start, num_cycles, width=2.8):
    for c in range(col_start, col_start + num_cycles):
        ws.column_dimensions[get_column_letter(c)].width = width

wb = Workbook()
ws = wb.active
ws.title = "Timing"

waves, logic = ld.load_timing_excel("input.xlsx")
waves_all = sim.simulate(waves, logic)

start_row = 2

groups = defaultdict(list)

for sig, data in waves_all.items():
    inst, local = split_hier(sig)
    groups[inst].append((sig, data))

row = start_row
first = True

for inst, sigs in groups.items():
    inst_start = row

    # 見出し行
    ws.cell(row=row, column=1, value=f"[{inst}]")
    row += 1

    for sig, (bw, values) in sigs:
        ws.cell(row=row, column=3, value="  " + sig)

        if first:
            set_wave_column_width(ws, 6, len(values), width=2.8)
            first = False

        dr.draw_wave(ws, row, 6, values, bw != 1)

        # --- 空行を詰める ---
        empty_row = row + 1
        ws.row_dimensions[empty_row].height = EMPTY_ROW_HEIGHT

        row += 2

    inst_end = row - 1

    ws.row_dimensions.group(
        inst_start + 1,
        inst_end,
        outline_level=1,
        hidden=True   # ← 初期状態で折りたたむ
    )

wb.save("timing_out.xlsx")

